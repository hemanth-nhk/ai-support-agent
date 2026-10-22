"""
Document ingestion, validation, extraction, chunking, and isolated Chroma storage.
Supports PDF, DOCX, and TXT files.
"""
from pathlib import Path
import io
import re
import uuid
from typing import Optional
import chromadb

from app.config import settings
from app.rag.embeddings import get_embedding_function

DOCUMENTS_COLLECTION_NAME = "uploaded_documents"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
DEFAULT_CHUNK_SIZE = 500  # characters
DEFAULT_CHUNK_OVERLAP = 50  # characters


class DocumentValidationError(Exception):
    """Raised when an uploaded document fails validation."""
    pass


class DocumentNotFoundError(Exception):
    """Raised when a document is not found for deletion or retrieval."""
    pass


def sanitize_filename(filename: str) -> str:
    """Sanitizes filename, stripping directory traversal and dangerous characters."""
    raw_name = Path(filename).name
    cleaned = re.sub(r"[^\w.\-_]", "_", raw_name).strip()
    return cleaned or "uploaded_document.txt"


def validate_file(filename: str, content_bytes: bytes) -> str:
    """
    Validates file extension, non-empty content, and max file size.
    Returns sanitized filename. Raises DocumentValidationError on failure.
    """
    if not content_bytes or len(content_bytes) == 0:
        raise DocumentValidationError("Document is empty.")

    if len(content_bytes) > MAX_FILE_SIZE_BYTES:
        raise DocumentValidationError(
            f"File exceeds maximum allowed size of {MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB."
        )

    sanitized = sanitize_filename(filename)
    ext = Path(sanitized).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentValidationError(
            f"Unsupported file type '{ext}'. Allowed types are: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )

    return sanitized


def extract_text(filename: str, content_bytes: bytes) -> str:
    """
    Extracts text from PDF, DOCX, or TXT file bytes.
    Raises DocumentValidationError if text cannot be extracted or is empty.
    """
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        try:
            text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content_bytes.decode("latin-1")
            except Exception as e:
                raise DocumentValidationError(f"Failed to decode text file: {e}")
        text = text.strip()
        if not text:
            raise DocumentValidationError("Document contains no readable text.")
        return text

    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content_bytes))
            pages_text = []
            for page in reader.pages:
                page_str = page.extract_text() or ""
                if page_str.strip():
                    pages_text.append(page_str.strip())
            text = "\n\n".join(pages_text).strip()
        except Exception as e:
            raise DocumentValidationError(f"Failed to parse PDF document: {e}")

        if not text:
            raise DocumentValidationError("Unable to extract text from PDF (document contains no readable text).")
        return text

    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            extracted_parts: list[str] = []

            # 1. Section Headers (deduplicate identical header text across sections)
            seen_headers: set[str] = set()
            for section in doc.sections:
                try:
                    header = section.header
                    if header:
                        for p in header.paragraphs:
                            t = p.text.strip()
                            if t and t not in seen_headers:
                                seen_headers.add(t)
                                extracted_parts.append(t)
                except Exception:
                    pass

            # 2. Body Paragraphs
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    extracted_parts.append(t)

            # 3. Table Rows and Cells (preserves readable ordering and cell content)
            for table in doc.tables:
                for row in table.rows:
                    row_cells: list[str] = []
                    seen_cells: set[int] = set()
                    for cell in row.cells:
                        cell_id = id(cell._tc)
                        if cell_id in seen_cells:
                            continue
                        seen_cells.add(cell_id)
                        cell_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                        if not cell_text and cell.text.strip():
                            cell_text = cell.text.strip()
                        if cell_text:
                            row_cells.append(cell_text)
                    if row_cells:
                        extracted_parts.append(" | ".join(row_cells))

            # 4. Section Footers (deduplicate identical footer text across sections)
            seen_footers: set[str] = set()
            for section in doc.sections:
                try:
                    footer = section.footer
                    if footer:
                        for p in footer.paragraphs:
                            t = p.text.strip()
                            if t and t not in seen_footers:
                                seen_footers.add(t)
                                extracted_parts.append(t)
                except Exception:
                    pass

            text = "\n".join(extracted_parts).strip()
        except Exception as e:
            raise DocumentValidationError(f"Failed to parse DOCX document: {e}")

        if not text:
            raise DocumentValidationError("Unable to extract text from DOCX (document contains no readable text).")
        return text

    else:
        raise DocumentValidationError(f"Unsupported file type '{ext}'.")


def normalize_text(text: str) -> str:
    """Normalizes whitespace and line breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """
    Deterministic chunking with sliding window and overlap.
    Suitable for embedding and RAG retrieval.
    """
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    step = max(1, chunk_size - chunk_overlap)
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start += step

    return chunks


def get_documents_collection(persist_dir: str | None = None, embedding_fn=None):
    """
    Retrieves or creates the isolated uploaded_documents collection in Chroma.
    Configured explicitly with cosine distance.
    """
    client = chromadb.PersistentClient(path=persist_dir or settings.chroma_persist_dir)
    emb_fn = embedding_fn or get_embedding_function()
    collection = client.get_or_create_collection(
        name=DOCUMENTS_COLLECTION_NAME,
        embedding_function=emb_fn,
        metadata={"hnsw:space": "cosine"},
    )
    return collection


def index_document(
    session_id: str,
    filename: str,
    content_bytes: bytes,
    persist_dir: str | None = None,
    embedding_fn=None,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> dict:
    """
    End-to-end ingestion pipeline:
    validates -> extracts -> normalizes -> chunks -> generates document_id -> stores in Chroma.
    """
    if not session_id or not session_id.strip():
        raise DocumentValidationError("A valid session_id is required.")

    sanitized_filename = validate_file(filename, content_bytes)
    raw_text = extract_text(sanitized_filename, content_bytes)
    normalized = normalize_text(raw_text)
    chunks = chunk_text(normalized, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    if not chunks:
        raise DocumentValidationError("Document produced no readable text chunks.")

    document_id = str(uuid.uuid4())
    collection = get_documents_collection(persist_dir=persist_dir, embedding_fn=embedding_fn)

    ids = [f"{document_id}_{i}" for i in range(len(chunks))]
    metadatas = [
        {
            "session_id": session_id.strip(),
            "document_id": document_id,
            "filename": sanitized_filename,
            "chunk_index": i,
        }
        for i in range(len(chunks))
    ]

    collection.add(
        ids=ids,
        documents=chunks,
        metadatas=metadatas,
    )

    return {
        "document_id": document_id,
        "filename": sanitized_filename,
        "chunk_count": len(chunks),
    }


def delete_document(
    document_id: str,
    session_id: str,
    persist_dir: str | None = None,
    embedding_fn=None,
) -> dict:
    """
    Deletes vectors/chunks for document_id from the uploaded_documents collection.
    Enforces non-empty session ownership.
    """
    if not document_id or not document_id.strip():
        raise DocumentValidationError("A valid document_id is required for deletion.")
    if not session_id or not session_id.strip():
        raise DocumentValidationError("A valid, non-empty session_id is required for document deletion.")

    clean_doc_id = document_id.strip()
    clean_session_id = session_id.strip()
    collection = get_documents_collection(persist_dir=persist_dir, embedding_fn=embedding_fn)

    where_filter = {
        "$and": [
            {"document_id": clean_doc_id},
            {"session_id": clean_session_id},
        ]
    }

    matches = collection.get(where=where_filter)
    matched_ids = matches.get("ids", [])

    if not matched_ids:
        raise DocumentNotFoundError(f"Document '{clean_doc_id}' not found.")

    collection.delete(ids=matched_ids)

    return {
        "status": "deleted",
        "document_id": clean_doc_id,
        "deleted_chunks": len(matched_ids),
    }


def get_session_documents(
    session_id: str,
    persist_dir: str | None = None,
    embedding_fn=None,
) -> list[dict]:
    """
    Returns unique documents ({document_id, filename}) uploaded in this session.
    """
    if not session_id or not session_id.strip():
        return []

    collection = get_documents_collection(persist_dir=persist_dir, embedding_fn=embedding_fn)
    if collection.count() == 0:
        return []

    clean_session_id = session_id.strip()
    matches = collection.get(where={"session_id": clean_session_id})
    metadatas = matches.get("metadatas", [])
    seen = {}
    for meta in metadatas:
        doc_id = meta.get("document_id")
        if doc_id and doc_id not in seen:
            seen[doc_id] = {
                "document_id": doc_id,
                "filename": meta.get("filename", "unknown"),
            }
    return list(seen.values())


def retrieve_document_chunks(
    session_id: str,
    query: str,
    document_id: str | None = None,
    persist_dir: str | None = None,
    embedding_fn=None,
    top_k: int = 3,
    threshold: float | None = None,
) -> list[dict]:
    """
    Retrieves document chunks strictly isolated by session_id.
    Optionally filters by specific document_id.
    Optionally filters by cosine distance threshold.
    Requires non-empty session_id.
    """
    if not session_id or not session_id.strip():
        raise DocumentValidationError("A valid session_id is required for document retrieval.")

    if not query or not query.strip():
        return []

    collection = get_documents_collection(persist_dir=persist_dir, embedding_fn=embedding_fn)
    if collection.count() == 0:
        return []

    clean_session_id = session_id.strip()
    where_filter: dict
    if document_id and document_id.strip():
        where_filter = {
            "$and": [
                {"session_id": clean_session_id},
                {"document_id": document_id.strip()},
            ]
        }
    else:
        where_filter = {"session_id": clean_session_id}

    effective_k = min(top_k, collection.count())
    if effective_k <= 0:
        return []

    results = collection.query(
        query_texts=[query.strip()],
        n_results=effective_k,
        where=where_filter,
    )

    docs = results.get("documents", [[]])[0] if results.get("documents") else []
    metas = results.get("metadatas", [[]])[0] if results.get("metadatas") else []
    dists = results.get("distances", [[]])[0] if results.get("distances") else []

    retrieved = []
    for doc, meta, dist in zip(docs, metas, dists):
        if threshold is not None and dist > threshold:
            continue
        retrieved.append({
            "text": doc,
            "metadata": meta,
            "distance": dist,
        })
    return retrieved
