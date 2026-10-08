"""
Benchmark Suite: Extraction Fidelity (Category B).
Measures recovered facts / total expected facts across TXT, PDF, and DOCX formats.
Identifies parser boundaries (e.g. DOCX tables, OCR-less scanned PDFs).
"""
import time
import io
import docx
import pytest

from app.services.document_service import extract_text, DocumentValidationError
from evaluation.reliability.models import BenchmarkCaseResult
from evaluation.reliability.fixtures_generator import (
    build_raw_pdf,
    build_scanned_pdf,
    build_docx_with_paragraphs,
    build_docx_with_headings_and_lists,
    build_docx_with_table,
    build_docx_mixed,
)

EXPECTED_FACTS = ["C1024", "30 days", "12 months", "50,000"]


def count_recovered_facts(text: str, facts: list[str]) -> tuple[int, list[str]]:
    lowered = text.lower()
    recovered = []
    for f in facts:
        # Check normalized form
        clean_f = f.lower().replace(",", "").replace("₹", "").replace("$", "")
        clean_text = lowered.replace(",", "").replace("₹", "").replace("$", "")
        if clean_f in clean_text:
            recovered.append(f)
    return len(recovered), recovered


def run_cases() -> list[BenchmarkCaseResult]:
    results: list[BenchmarkCaseResult] = []

    # Case EXT_01: Plain TXT
    t0 = time.perf_counter()
    txt_content = (
        "Customer Service Summary:\n"
        "Customer ID: C1024\n"
        "Refund period: 30 days\n"
        "Warranty: 12 months\n"
        "Approval limit: ₹50,000\n"
    ).encode("utf-8")
    extracted = extract_text("facts.txt", txt_content)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_01",
        category="EXTRACTION",
        file_type="TXT",
        scenario="Plain UTF-8 TXT extraction with known facts",
        expected_behavior=f"Recover all 4 facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_02: Standard single-page PDF
    t0 = time.perf_counter()
    pdf_bytes = build_raw_pdf([
        "Customer Service Summary:\nCustomer ID: C1024\nRefund period: 30 days\nWarranty: 12 months\nApproval limit: 50,000"
    ])
    extracted = extract_text("facts.pdf", pdf_bytes)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_02",
        category="EXTRACTION",
        file_type="PDF",
        scenario="Single-page text PDF extraction",
        expected_behavior=f"Recover all 4 facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_03: Multi-page PDF
    t0 = time.perf_counter()
    pdf_mp = build_raw_pdf([
        "Page 1: Account Overview\nCustomer ID: C1024",
        "Page 2: Operations & Policies\nRefund period: 30 days",
        "Page 3: Hardware Support\nWarranty: 12 months\nApproval limit: 50,000",
    ])
    extracted = extract_text("facts_multipage.pdf", pdf_mp)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_03",
        category="EXTRACTION",
        file_type="PDF",
        scenario="Multi-page PDF spanning 3 pages",
        expected_behavior=f"Recover facts across all pages ({len(EXPECTED_FACTS)}/4)",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_04: PDF with table-formatted text
    t0 = time.perf_counter()
    pdf_table_txt = (
        "Item               Value\n"
        "Customer ID        C1024\n"
        "Refund period      30 days\n"
        "Warranty           12 months\n"
        "Approval limit     50,000\n"
    )
    pdf_table = build_raw_pdf([pdf_table_txt])
    extracted = extract_text("facts_table.pdf", pdf_table)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_04",
        category="EXTRACTION",
        file_type="PDF",
        scenario="Text table in PDF layout",
        expected_behavior=f"Recover all 4 tabular facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P2",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_05: PDF with repeated header/footer
    t0 = time.perf_counter()
    p1 = "CONFIDENTIAL HEADER\nCustomer ID: C1024\nCONFIDENTIAL FOOTER"
    p2 = "CONFIDENTIAL HEADER\nRefund period: 30 days\nCONFIDENTIAL FOOTER"
    p3 = "CONFIDENTIAL HEADER\nWarranty: 12 months\nApproval limit: 50,000\nCONFIDENTIAL FOOTER"
    pdf_hf = build_raw_pdf([p1, p2, p3])
    extracted = extract_text("facts_hf.pdf", pdf_hf)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_05",
        category="EXTRACTION",
        file_type="PDF",
        scenario="PDF with repeated headers and footers across pages",
        expected_behavior="Recover all 4 facts despite header/footer repetition",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P2",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_06: PDF with complex multi-column layout
    t0 = time.perf_counter()
    col_text = (
        "Column Left                     Column Right\n"
        "Customer ID: C1024              Warranty: 12 months\n"
        "Refund period: 30 days          Approval limit: 50,000\n"
    )
    pdf_cols = build_raw_pdf([col_text])
    extracted = extract_text("facts_cols.pdf", pdf_cols)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_06",
        category="EXTRACTION",
        file_type="PDF",
        scenario="PDF with multi-column text layout",
        expected_behavior="Recover facts across columnar presentation",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P2",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_07: Scanned / image-only PDF (No OCR)
    t0 = time.perf_counter()
    scanned_bytes = build_scanned_pdf()
    try:
        extract_text("scanned_doc.pdf", scanned_bytes)
        passed = False
        act = "Unexpectedly extracted text from image-only PDF"
        f_class = "EXTRACTION"
        sev = "P2"
    except DocumentValidationError as e:
        # Explicit capability boundary: NO OCR
        passed = True
        act = f"Fails closed cleanly with DocumentValidationError: {e} (Classified as UNSUPPORTED_CAPABILITY: No OCR)"
        f_class = "UNSUPPORTED_CAPABILITY"
        sev = "P2"
    except Exception as e:
        passed = False
        act = f"Unexpected error: {e}"
        f_class = "EXTRACTION"
        sev = "P1"
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_07",
        category="EXTRACTION",
        file_type="PDF",
        scenario="Scanned/image-only PDF without embedded text streams",
        expected_behavior="Explicit capability boundary rejection: No readable text (No OCR)",
        actual_behavior=act,
        passed=passed,
        failure_class=f_class,
        severity=sev,
        latency_ms=round(lat, 2),
        notes="Documented architecture boundary: Local offline text extractor deliberately lacks heavy OCR engine."
    ))

    # Case EXT_08: DOCX standard paragraphs
    t0 = time.perf_counter()
    docx_paras = [
        "Customer Service Record",
        "Customer ID: C1024",
        "Refund period: 30 days",
        "Warranty: 12 months",
        "Approval limit: 50,000",
    ]
    extracted = extract_text("facts.docx", build_docx_with_paragraphs(docx_paras))
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_08",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="Standard DOCX paragraphs",
        expected_behavior=f"Recover all 4 facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_09: DOCX headings + body text
    t0 = time.perf_counter()
    h_doc = docx.Document()
    h_doc.add_heading("Customer ID: C1024", level=1)
    h_doc.add_paragraph("Policy terms stipulate a refund period: 30 days.")
    h_doc.add_heading("Hardware Coverage", level=2)
    h_doc.add_paragraph("Warranty: 12 months.")
    h_doc.add_paragraph("Approval limit: 50,000.")
    buf = io.BytesIO()
    h_doc.save(buf)
    extracted = extract_text("facts_headings.docx", buf.getvalue())
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_09",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX headings mixed with body paragraphs",
        expected_behavior=f"Recover all 4 facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_10: DOCX bulleted lists
    t0 = time.perf_counter()
    items = ["Customer ID: C1024", "Refund period: 30 days", "Warranty: 12 months", "Approval limit: 50,000"]
    docx_lists = build_docx_with_headings_and_lists(["Key Fact Sheet"], items)
    extracted = extract_text("facts_list.docx", docx_lists)
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_10",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX bulleted list items",
        expected_behavior=f"Recover all 4 list facts: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_11: DOCX table (All 4 facts in table cells)
    # Parser boundary discovery: python-docx extract_text currently reads doc.paragraphs, not doc.tables!
    t0 = time.perf_counter()
    headers = ["Parameter", "Value"]
    rows = [
        ["Customer ID", "C1024"],
        ["Refund period", "30 days"],
        ["Warranty", "12 months"],
        ["Approval limit", "50,000"],
    ]
    # Add a minimal paragraph so docx doesn't fail empty validation immediately
    table_doc = docx.Document()
    table_doc.add_paragraph("Table of Parameters")
    t_obj = table_doc.add_table(rows=len(rows) + 1, cols=2)
    t_obj.cell(0, 0).text = headers[0]
    t_obj.cell(0, 1).text = headers[1]
    for r_idx, r in enumerate(rows):
        t_obj.cell(r_idx + 1, 0).text = r[0]
        t_obj.cell(r_idx + 1, 1).text = r[1]
    buf = io.BytesIO()
    table_doc.save(buf)
    extracted = extract_text("facts_table.docx", buf.getvalue())
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    # Record actual behavior without modifying production code
    passed = (rec_count == 4)
    actual_behavior = f"Recovered {rec_count}/4 facts: {rec_list}" if passed else f"Recovered {rec_count}/4 facts ({rec_list}). Extractor parsed: '{extracted}'"
    notes = None if passed else "Documented parser limitation: extract_text() in app/services/document_service.py iterates doc.paragraphs only, missing doc.tables."
    results.append(BenchmarkCaseResult(
        case_id="EXT_11",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX table structure containing all target facts",
        expected_behavior="Recover table cell facts",
        actual_behavior=actual_behavior,
        passed=passed,
        failure_class="EXTRACTION" if not passed else None,
        severity="P1" if not passed else None,
        latency_ms=round(lat, 2),
        notes=notes,
    ))

    # Case EXT_12: DOCX multiple sections
    t0 = time.perf_counter()
    sec_doc = docx.Document()
    sec_doc.add_paragraph("Section 1: Customer ID: C1024")
    sec_doc.add_section()
    sec_doc.add_paragraph("Section 2: Refund period: 30 days")
    sec_doc.add_section()
    sec_doc.add_paragraph("Section 3: Warranty: 12 months and Approval limit: 50,000")
    buf = io.BytesIO()
    sec_doc.save(buf)
    extracted = extract_text("facts_sections.docx", buf.getvalue())
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_12",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX with multiple section breaks",
        expected_behavior=f"Recover all 4 facts across sections: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P1",
        latency_ms=round(lat, 2)
    ))

    # Case EXT_13: DOCX header / footer
    t0 = time.perf_counter()
    hf_doc = docx.Document()
    sec = hf_doc.sections[0]
    sec.header.paragraphs[0].text = "Header: Customer ID: C1024"
    sec.footer.paragraphs[0].text = "Footer: Approval limit: 50,000"
    hf_doc.add_paragraph("Body: Refund period: 30 days. Warranty: 12 months.")
    buf = io.BytesIO()
    hf_doc.save(buf)
    extracted = extract_text("facts_header_footer.docx", buf.getvalue())
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    passed = (rec_count == 4)
    results.append(BenchmarkCaseResult(
        case_id="EXT_13",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX with facts split between headers/footers and body",
        expected_behavior="Recover facts including header/footer annotations",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=passed,
        failure_class="EXTRACTION" if not passed else None,
        severity="P2" if not passed else None,
        latency_ms=round(lat, 2),
        notes=None if passed else "Documented limitation: python-docx body paragraphs do not include header/footer elements unless specifically traversed."
    ))

    # Case EXT_14: DOCX mixed (Paragraphs + Table)
    t0 = time.perf_counter()
    mixed_paras = ["Customer ID: C1024", "Refund period: 30 days"]
    mixed_table = [["Warranty", "12 months"], ["Approval limit", "50,000"]]
    extracted = extract_text("facts_mixed.docx", build_docx_mixed(mixed_paras, mixed_table))
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    passed = (rec_count == 4)
    results.append(BenchmarkCaseResult(
        case_id="EXT_14",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX with 2 facts in paragraphs and 2 facts in table",
        expected_behavior="Recover all 4 facts across paragraph and table sections",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=passed,
        failure_class="EXTRACTION" if not passed else None,
        severity="P1" if not passed else None,
        latency_ms=round(lat, 2)
    ))

    # Case EXT_15: DOCX with surrounding filler noise
    t0 = time.perf_counter()
    noise = "This is filler context explaining operational guidelines in extensive length without any specific parameters."
    noisy_paras = [noise] * 5 + [
        "Customer ID: C1024",
        "Refund period: 30 days",
        "Warranty: 12 months",
        "Approval limit: 50,000",
    ] + [noise] * 5
    extracted = extract_text("facts_noisy.docx", build_docx_with_paragraphs(noisy_paras))
    rec_count, rec_list = count_recovered_facts(extracted, EXPECTED_FACTS)
    lat = (time.perf_counter() - t0) * 1000
    results.append(BenchmarkCaseResult(
        case_id="EXT_15",
        category="EXTRACTION",
        file_type="DOCX",
        scenario="DOCX with facts embedded amidst extensive filler paragraphs",
        expected_behavior=f"Recover all 4 facts despite noise: {EXPECTED_FACTS}",
        actual_behavior=f"Recovered {rec_count}/4 facts: {rec_list}",
        passed=(rec_count == 4),
        failure_class=None if rec_count == 4 else "EXTRACTION",
        severity=None if rec_count == 4 else "P2",
        latency_ms=round(lat, 2)
    ))

    return results


def test_extraction_suite():
    cases = run_cases()
    # At least the core text, PDF, and standard DOCX paragraph extractions must pass
    assert any(c.case_id == "EXT_01" and c.passed for c in cases)
    assert any(c.case_id == "EXT_02" and c.passed for c in cases)
    assert any(c.case_id == "EXT_08" and c.passed for c in cases)
